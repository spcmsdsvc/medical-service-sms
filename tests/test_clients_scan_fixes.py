"""Client List scan fixes (plans.md "Client List — scan fixes and table improvements")."""
import datetime
import io
import pathlib
import re
import unittest

from tests.test_tsr_autosave_client_groups import APP_IMPORT_ERROR, ClientGroupRouteTests, app_module

ROOT = pathlib.Path(__file__).resolve().parents[1]


def css_rule(source, selector):
    match = re.search(re.escape(selector) + r'\s*\{([^}]*)\}', source)
    return match.group(1) if match else ''


class ClientListSourceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.page = (ROOT / 'templates' / 'clients.html').read_text(encoding='utf-8')

    def test_batch1_admin_flag_and_dead_code(self):
        self.assertNotIn("current_user.role in ['superadmin', 'regional_admin']", self.page)
        self.assertIn("const isAdminUser = {{ 'true' if can_admin_clients else 'false' }};", self.page)
        self.assertNotIn('clientHasDesignationMatch', self.page)
        self.assertIn("clientAlert('Allow pop-ups to print this report.')", self.page)

    def test_batch2_keyboard_and_labels(self):
        render_table = self.page.split('function renderTable(data) {', 1)[1].split('\nfunction ', 1)[0]
        mobile = self.page.split('function renderClientMobileCards(data){', 1)[1].split('\nfunction ', 1)[0]
        self.assertIn('<button type="button" class="client-name-link"', render_table)
        self.assertIn('client-more-link', render_table)
        self.assertNotIn('badge bg-info', self.page)
        self.assertNotIn('btn-outline-dark', render_table)
        self.assertIn('client-name-link client-mobile-name', mobile)
        self.assertIn("btn.closest('th')?.setAttribute('aria-sort'", self.page)
        self.assertIn('aria-labelledby="view-client-name"', self.page)
        self.assertIn("el.setAttribute('aria-expanded'", self.page)
        self.assertNotIn('outline: none', css_rule(self.page, '.client-sort-btn:focus'))

    def test_batch3_phone_cards_follow_theme(self):
        for selector in ('.client-mobile-card', '.client-mobile-name', '.client-mobile-address',
                         '.client-mobile-contact-box', '.client-mobile-contact-title', '.client-mobile-contact-lines',
                         '.client-mobile-view', '.client-mobile-empty'):
            rule = css_rule(self.page, selector)
            self.assertTrue(rule, selector)
            self.assertNotRegex(rule, r'#(0f172a|64748b|e2e8f0|f8fafc|fff\b|cbd5e1)', selector)
        self.assertNotIn('#bbf7d0', css_rule(self.page, '.client-sort-btn.active'))
        self.assertNotIn('#86efac', self.page)
        self.assertIn('width: 44px', css_rule(self.page, '.client-mobile-icon-action'))
        self.assertIn('border-radius: 16px', css_rule(self.page, '.client-mobile-card'))
        self.assertNotIn('client-mobile-icon-action disabled', self.page)

    def test_batch4_table_layout(self):
        self.assertIn("const CLIENT_SORT_KEYS = ['name', 'address', 'group', 'contact', 'equipment'];", self.page)
        self.assertIn('<th class="client-col-group">', self.page)
        self.assertIn('.client-table.is-hidden-group .client-col-group', self.page)
        self.assertIn('position: sticky', css_rule(self.page, '.client-table thead th'))
        self.assertLess(self.page.index('id="client-count-label"'), self.page.index('id="client-table"'))
        view = self.page.split('async function viewContacts(id){', 1)[1].split('\nasync function ', 1)[0]
        self.assertLess(view.index("'products-sec'"), view.index("'contacts-sec'"))
        self.assertIn('client-status-chip', view)


@unittest.skipIf(app_module is None, f'app import failed: {APP_IMPORT_ERROR}')
class ClientListRouteTests(unittest.TestCase):
    isolated_database = ClientGroupRouteTests.isolated_database

    def _make_client(self, name, with_undated_product=False):
        with app_module.app.app_context():
            record = app_module.Client(name=name, address='Somewhere', group_name='North')
            app_module.db.session.add(record)
            app_module.db.session.flush()
            if with_undated_product:
                app_module.db.session.add(app_module.Product(
                    serial_number=f'SN-{record.id}', name='Undated Unit', client_id=record.id,
                ))
                app_module.db.session.add(app_module.Product(
                    serial_number=f'SN-{record.id}-B', name='Dated Unit', client_id=record.id,
                    start_warranty_date=datetime.date(2025, 1, 2), end_warranty_date=datetime.date(2027, 1, 2),
                ))
            app_module.db.session.commit()
            return record.id

    def test_excel_filename_is_quoted_and_dates_blank(self):
        from openpyxl import load_workbook
        with self.isolated_database() as clients:
            client_id = self._make_client('Allied Care, Inc. Biñan', with_undated_product=True)
            response = clients['admin'].get(f'/export_client_excel/{client_id}')
            self.assertEqual(response.status_code, 200)
            disposition = response.headers['Content-Disposition']
            self.assertIn('attachment', disposition)
            self.assertIn("filename*=UTF-8''", disposition)
            self.assertNotRegex(disposition, r'filename=[^";]*,')
            rows = [[cell for cell in row] for row in load_workbook(io.BytesIO(response.data)).active.iter_rows(values_only=True)]
            flat = [cell for row in rows for cell in row]
            self.assertNotIn('None', flat)
            self.assertIn(['Group', 'North'], [list(row[:2]) for row in rows])
            self.assertIn('2025-01-02', flat)

    def test_admin_buttons_follow_server_admin_check(self):
        with self.isolated_database() as clients:
            with app_module.app.app_context():
                rogue = app_module.User(username='client-rogue-superadmin', password='test-only',
                                        role='staff', is_active=True)
                app_module.db.session.add(rogue)
                app_module.db.session.commit()
                rogue_id = rogue.id
            rogue_client = app_module.app.test_client()
            with rogue_client.session_transaction() as session:
                session['_user_id'] = str(rogue_id)
                session['_fresh'] = True
            page = rogue_client.get('/clients_page').get_data(as_text=True)
            self.assertNotIn('/export_clients', page)
            self.assertIn("const isAdminUser = false;", page)
            admin_page = clients['admin'].get('/clients_page').get_data(as_text=True)
            self.assertIn('/export_clients', admin_page)
            self.assertIn("const isAdminUser = true;", admin_page)

    def test_check_client_duplicate(self):
        with self.isolated_database() as clients:
            client_id = self._make_client('Duplicate Check Hospital')
            match = clients['admin'].get('/check_client_duplicate?name=Duplicate Check Hospital&address=Somewhere').get_json()
            self.assertTrue(match['match'])
            self.assertEqual(match['id'], client_id)
            excluded = clients['admin'].get(
                f'/check_client_duplicate?name=Duplicate Check Hospital&address=Somewhere&exclude_id={client_id}'
            ).get_json()
            self.assertFalse(excluded['match'])
            self.assertFalse(clients['admin'].get('/check_client_duplicate?name=').get_json()['match'])


if __name__ == '__main__':
    unittest.main()
