"""P.O. Details scan batches (plans.md, 2026-10-26): edit-mode machines, delete confirm,
dark mode, header and export fixes."""

import pathlib
import re
import unittest
import uuid
from datetime import date
from io import BytesIO

import app as app_module  # noqa: E402
from openpyxl import load_workbook


ROOT = pathlib.Path(__file__).resolve().parents[1]
TEMPLATE = (ROOT / 'templates' / 'po_details.html').read_text(encoding='utf-8')
STYLE = TEMPLATE.split('</style>', 1)[0]


class PoDetailsTemplateTests(unittest.TestCase):
    """Source checks for the page-only parts of each batch."""

    def test_batch1_edit_mode_ignores_renewal_flags(self):
        self.assertIn('const renewal = !state.editingId && productIsRenewalEligible(product);', TEMPLATE)
        self.assertIn('renewal !== (!state.editingId && productIsRenewalEligible(existing))', TEMPLATE)
        self.assertIn('if (!state.editingId && productIsRenewalEligible(product)) {', TEMPLATE)

    def test_batch2_danger_confirm_opens_on_cancel_traps_tab_and_restores_focus(self):
        self.assertIn('(danger ? cancelButton : confirmButton).focus()', TEMPLATE)
        self.assertIn("if (event.key === 'Tab')", TEMPLATE)
        self.assertIn('const opener = document.activeElement;', TEMPLATE)
        self.assertIn('opener.focus();', TEMPLATE)

    def test_batch3_tokens_badges_weights_and_buttons(self):
        self.assertNotIn('--app-table-head', TEMPLATE)
        self.assertNotIn('--app-focus-ring', TEMPLATE)
        for badge in ('single', 'semi_annual', 'quarterly'):
            self.assertIn(f'[data-app-theme="dark"] .po-type-badge.{badge} {{', STYLE)
        self.assertEqual(re.findall(r'font-weight:\s*([7-9]\d\d)', STYLE), [])
        self.assertIn('class="btn btn-outline-primary btn-sm" id="po-export-button"', TEMPLATE)
        self.assertIn(':root[data-app-theme="dark"] .po-page .btn-outline-danger', STYLE)

    def test_batch4_page_header_stat_cards_and_states(self):
        self.assertIn('<h1 class="dash-page-h1">P.O. Details</h1>', TEMPLATE)
        self.assertIn('class="btn btn-success" id="po-add-button"', TEMPLATE)
        self.assertNotIn('po-hero', TEMPLATE)
        self.assertIn("css/app-dashboard.css", TEMPLATE)
        for element_id in (
            'po-total-count', 'po-single-count', 'po-single-total', 'po-semi-annual-count',
            'po-semi-annual-total', 'po-quarterly-count', 'po-quarterly-total', 'po-client-count',
        ):
            self.assertIn(f'id="{element_id}"', TEMPLATE)
        self.assertIn('aria-busy="true" id="po-stats"', TEMPLATE)
        self.assertIn("setAttribute('aria-sort'", TEMPLATE)
        self.assertIn('id="po-retry-button"', TEMPLATE)
        self.assertIn('id="po-filters-toggle" aria-expanded="false" aria-controls="po-filter-panel"', TEMPLATE)
        self.assertIn('.po-filter-panel:not(.is-open)', STYLE)
        self.assertIn("response.headers.get('Content-Disposition')", TEMPLATE)
        self.assertNotIn('Warranty ends', TEMPLATE)


class PoDetailsEndpointTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = app_module.app
        cls.app.config.update(TESTING=True, WTF_CSRF_ENABLED=False)
        cls.suffix = uuid.uuid4().hex[:8]
        with cls.app.app_context():
            app_module.db.create_all()
            app_module.ensure_user_admin_capability_columns()
            app_module.ensure_purchase_order_schema()
            user = app_module.User(
                username=f'po_batches_{cls.suffix}',
                password=app_module.generate_password_hash(uuid.uuid4().hex),
                role='staff', is_active=True, po_admin_access=True,
            )
            client = app_module.Client(name=f'PO Batches Client {cls.suffix}', address='Address')
            app_module.db.session.add_all([user, client])
            app_module.db.session.flush()
            cls.serials = [f'PO-BATCH-A-{cls.suffix}', f'PO-BATCH-B-{cls.suffix}', f'PO-BATCH-C-{cls.suffix}']
            # A and B: coverage that has ended, matching dates. C: current coverage for the export.
            for serial, start, end in (
                (cls.serials[0], date(2020, 1, 1), date(2022, 12, 31)),
                (cls.serials[1], date(2020, 1, 1), date(2022, 12, 31)),
                (cls.serials[2], date(2026, 8, 25), date(2028, 8, 25)),
            ):
                app_module.db.session.add(app_module.Product(
                    serial_number=serial, name='Batch Machine', client_id=client.id,
                    start_warranty_date=start, end_warranty_date=end, under_contract=True,
                ))
            app_module.db.session.commit()
            cls.user_id = user.id
            cls.client_id = client.id

    @classmethod
    def tearDownClass(cls):
        with cls.app.app_context():
            for po in app_module.PurchaseOrder.query.filter_by(client_id=cls.client_id).all():
                app_module.db.session.delete(po)
            app_module.db.session.flush()
            for serial in cls.serials:
                product = app_module.db.session.get(app_module.Product, serial)
                if product:
                    app_module.db.session.delete(product)
            for model, key in ((app_module.Client, cls.client_id), (app_module.User, cls.user_id)):
                record = app_module.db.session.get(model, key)
                if record:
                    app_module.db.session.delete(record)
            app_module.db.session.commit()

    def setUp(self):
        # Session login, as test_purchase_orders does: repeated /login posts hit the rate limit.
        self.client = self.app.test_client()
        with self.client.session_transaction() as session:
            session['_user_id'] = str(self.user_id)
            session['_fresh'] = True

    def _create(self, number, serials, po_type, amount):
        response = self.client.post('/add_purchase_order', json={
            'client_id': self.client_id, 'product_serials': serials, 'po_number': number,
            'po_type': po_type, 'amount': amount, 'date_mode': 'product',
        })
        self.assertEqual(response.status_code, 201, response.get_data(as_text=True))
        return response.get_json()['purchase_order']

    def test_edit_can_add_a_machine_whose_latest_po_has_ended(self):
        first = self._create(f'EDIT-A-{self.suffix}', [self.serials[0]], 'quarterly', '1000')
        self._create(f'EDIT-B-{self.suffix}', [self.serials[1]], 'quarterly', '1000')
        response = self.client.put(f"/update_purchase_order/{first['id']}", json={
            'client_id': self.client_id, 'product_serials': self.serials[:2],
            'po_number': f'EDIT-A-{self.suffix}', 'po_type': 'quarterly', 'amount': '1000',
            'date_mode': 'product',
        })
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        self.assertEqual(response.get_json()['purchase_order']['machine_serials'], self.serials[:2])

    def test_export_amount_sort_follows_po_amount_and_total_k_needs_a_date_filter(self):
        # Quarterly 1000 (125 per visit) vs annual Single 600 (200 per visit): in fiscal 2026 Q2
        # the filtered sums cross the P.O. amounts, so the old sort put the Single first.
        quarterly = f'EXP-Q-{self.suffix}'
        single = f'EXP-S-{self.suffix}'
        self._create(quarterly, [self.serials[2]], 'quarterly', '1000')
        self._create(single, [self.serials[2]], 'single', '600')
        query = f'client=PO Batches Client {self.suffix}&machine={self.serials[2]}&sort=amount&direction=desc'

        filtered = self.client.get(f'/export_purchase_orders?{query}&fiscal_year=2026&fiscal_quarter=2')
        self.assertEqual(filtered.status_code, 200)
        sheet = load_workbook(BytesIO(filtered.data)).active
        self.assertEqual([sheet.cell(row, 4).value for row in (2, 3)], [quarterly, single])
        self.assertTrue(str(sheet.cell(5, 11).value).startswith('=SUM(K2:K3)'))

        unfiltered = self.client.get(f'/export_purchase_orders?{query}')
        sheet = load_workbook(BytesIO(unfiltered.data)).active
        self.assertEqual(sheet.cell(5, 1).value, 'TOTAL AMOUNT')
        self.assertEqual(sheet.cell(5, 10).value, '=SUM(J2:J3)')
        self.assertIsNone(sheet.cell(5, 11).value)

    def test_products_payload_drops_duplicate_date_keys(self):
        products = self.client.get('/get_purchase_orders').get_json()['products']
        product = next(item for item in products if item['serial_number'] == self.serials[0])
        self.assertEqual((product['start_date'], product['end_date']), ('2020-01-01', '2022-12-31'))
        for key in ('product_start_date', 'product_end_date', 'start_warranty_date', 'end_warranty_date', 'end_warranty'):
            self.assertNotIn(key, product)


if __name__ == '__main__':
    unittest.main()
